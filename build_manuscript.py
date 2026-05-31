#!/usr/bin/env python3
"""Assemble an Elsevier-style submission manuscript (.docx + .md) from the
verified review content. Single column, double-spaced, continuous line
numbers, page numbers, title page, Highlights, structured front matter,
abbreviations, and declaration sections."""
import re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "review_chemical_processes_climate.md"
OUT_DOCX = "review_manuscript_elsevier.docx"
OUT_MD = "review_manuscript.md"

# ---------------------------------------------------------------- read source
with open(SRC, encoding="utf-8") as f:
    lines = f.read().split("\n")

def first_idx(pred):
    for i, l in enumerate(lines):
        if pred(l):
            return i
    return -1

title = next(l[2:].strip() for l in lines if l.startswith("# "))
i_abs = first_idx(lambda l: l.strip() == "## Abstract")
i_intro = first_idx(lambda l: l.startswith("## 1. Introduction"))
i_refs = first_idx(lambda l: l.strip() == "## References")

# abstract + keywords
abstract_parts, keywords = [], ""
for l in lines[i_abs + 1:i_intro]:
    s = l.strip()
    if not s or s == "---":
        continue
    if s.startswith("**Keywords:**"):
        keywords = s.replace("**Keywords:**", "").strip()
        continue
    abstract_parts.append(s)
abstract_text = " ".join(abstract_parts)

body_lines = lines[i_intro:i_refs]
ref_lines = lines[i_refs:]

# ---------------------------------------------------------------- build blocks
# block = (kind, text); kinds: title, authors, affil, corr, h1, h2, para,
# bullet, kw, ref, pagebreak, small
blocks = []

# ---- Title page
blocks.append(("title", title))
blocks.append(("authors", "[First Author]ᵃ,*, [Second Author]ᵇ, [Third Author]ᵃ"))
blocks.append(("affil", "ᵃ [Department, Institution, City, Postal code, Country]"))
blocks.append(("affil", "ᵇ [Department, Institution, City, Postal code, Country]"))
blocks.append(("corr", "* Corresponding author. E-mail address: [corresponding.author@institution.edu] ([First Author])."))
blocks.append(("corr", "ORCID: [0000-0000-0000-0000]"))
blocks.append(("pagebreak", ""))

# ---- Highlights (<=85 characters each)
highlights = [
    "Climate change is the macroscopic expression of coupled cross-sphere chemistry.",
    "Atmosphere, ocean and rock chemistry are linked by global biogeochemical cycles.",
    "Carbon–climate, permafrost and methane feedbacks amplify warming on decadal scales.",
    "The silicate-weathering negative feedback stabilizes climate only over ~10^5 years.",
    "Fast amplifying and slow stabilizing feedbacks act on mismatched timescales.",
]
blocks.append(("h1", "Highlights"))
for h in highlights:
    blocks.append(("bullet", h))

# ---- Abstract
blocks.append(("h1", "Abstract"))
blocks.append(("para", abstract_text))
blocks.append(("kw", keywords))

# ---- Abbreviations
abbrev = [
    ("AMOC", "Atlantic Meridional Overturning Circulation"),
    ("BC", "black carbon"),
    ("CCN", "cloud condensation nuclei"),
    ("CDR", "carbon dioxide removal"),
    ("DIC", "dissolved inorganic carbon"),
    ("DMS", "dimethyl sulfide"),
    ("ECS", "equilibrium climate sensitivity"),
    ("ERF", "effective radiative forcing"),
    ("ERFaci", "ERF from aerosol–cloud interactions"),
    ("ERFari", "ERF from aerosol–radiation interactions"),
    ("ERW", "enhanced rock weathering"),
    ("GWP", "global warming potential"),
    ("HNLC", "high-nitrate, low-chlorophyll"),
    ("IPCC AR6", "Intergovernmental Panel on Climate Change Sixth Assessment Report"),
    ("MAOM", "mineral-associated organic matter"),
    ("OH", "hydroxyl radical"),
    ("OMZ", "oxygen minimum zone"),
    ("POM", "particulate organic matter"),
    ("SOA", "secondary organic aerosol"),
    ("TCR", "transient climate response"),
]
blocks.append(("h1", "Abbreviations"))
for ab, full in abbrev:
    blocks.append(("abbr", f"{ab}\t{full}"))
blocks.append(("pagebreak", ""))

# ---- Main text (parse body)
def parse_content(src_lines, refs=False, stop_note=True):
    out = []
    for l in src_lines:
        s = l.strip()
        if not s or s == "---":
            continue
        if stop_note and s.startswith("*Prepared as"):
            break
        if s.startswith("## "):
            out.append(("h1", s[3:].strip()))
        elif s.startswith("### "):
            out.append(("h2", s[4:].strip()))
        elif refs and re.match(r"^\[\d+\]", s):
            out.append(("ref", s))
        else:
            out.append(("para", s))
    return out

blocks += parse_content(body_lines)

# ---- Declarations
blocks.append(("h1", "CRediT authorship contribution statement"))
blocks.append(("para", "[First Author]: Conceptualization, Investigation, Writing – original draft, Visualization. [Second Author]: Methodology, Writing – review & editing. [Third Author]: Supervision, Writing – review & editing."))
blocks.append(("h1", "Declaration of competing interest"))
blocks.append(("para", "The authors declare that they have no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper."))
blocks.append(("h1", "Funding"))
blocks.append(("para", "This research did not receive any specific grant from funding agencies in the public, commercial, or not-for-profit sectors. [Replace with grant numbers if applicable.]"))
blocks.append(("h1", "Data availability"))
blocks.append(("para", "No new data were generated or analysed in this review. All data discussed are available in the cited published sources."))
blocks.append(("h1", "Acknowledgements"))
blocks.append(("para", "[Acknowledge colleagues, reviewers, or institutional support here.]"))

# ---- References
blocks += parse_content(ref_lines, refs=True)

# ---------------------------------------------------------------- docx render
INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")

def add_runs(p, text, base_size=12):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*") and tok.endswith("*"):
            r = p.add_run(tok[1:-1]); r.italic = True
        else:
            p.add_run(tok)

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(12)
st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
st.paragraph_format.space_after = Pt(0)

def single(p):
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

for kind, text in blocks:
    if kind == "title":
        p = doc.add_paragraph(); single(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.bold = True; r.font.size = Pt(15)
        p.paragraph_format.space_after = Pt(18)
    elif kind == "authors":
        p = doc.add_paragraph(); single(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.font.size = Pt(12)
        p.paragraph_format.space_after = Pt(8)
    elif kind == "affil":
        p = doc.add_paragraph(); single(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.italic = True; r.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(2)
    elif kind == "corr":
        p = doc.add_paragraph(); single(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(2)
    elif kind == "pagebreak":
        doc.add_page_break()
    elif kind == "h1":
        p = doc.add_paragraph(); single(p)
        p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text); r.bold = True; r.font.size = Pt(13)
    elif kind == "h2":
        p = doc.add_paragraph(); single(p)
        p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text); r.bold = True; r.italic = True; r.font.size = Pt(12)
    elif kind == "bullet":
        p = doc.add_paragraph(); single(p)
        p.paragraph_format.left_indent = Pt(18); p.paragraph_format.space_after = Pt(2)
        p.add_run("•  " + text)
    elif kind == "kw":
        p = doc.add_paragraph(); single(p)
        p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(6)
        r = p.add_run("Keywords: "); r.bold = True
        p.add_run(text)
    elif kind == "abbr":
        ab, full = text.split("\t", 1)
        p = doc.add_paragraph(); single(p); p.paragraph_format.space_after = Pt(0)
        r = p.add_run(ab); r.bold = True
        p.add_run("\t" + full)
    elif kind == "ref":
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Pt(22)
        p.paragraph_format.first_line_indent = Pt(-22)
        add_runs(p, text)
    else:  # para
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.first_line_indent = Pt(18)
        add_runs(p, text)

# ---- continuous line numbers
sectPr = doc.sections[0]._sectPr
ln = OxmlElement("w:lnNumType")
ln.set(qn("w:countBy"), "1")
ln.set(qn("w:restart"), "continuous")
ln.set(qn("w:distance"), "360")
sectPr.append(ln)

# ---- page number in footer (centered)
footer = doc.sections[0].footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
run._r.append(b); run._r.append(instr); run._r.append(e)

doc.save(OUT_DOCX)

# ---------------------------------------------------------------- markdown mirror
md = []
for kind, text in blocks:
    if kind == "title":
        md.append(f"# {text}\n")
    elif kind in ("authors", "affil", "corr"):
        md.append(text + "  ")
    elif kind == "pagebreak":
        md.append("\n---\n")
    elif kind == "h1":
        md.append(f"\n## {text}\n")
    elif kind == "h2":
        md.append(f"\n### {text}\n")
    elif kind == "bullet":
        md.append(f"- {text}")
    elif kind == "kw":
        md.append(f"\n**Keywords:** {text}\n")
    elif kind == "abbr":
        ab, full = text.split("\t", 1)
        md.append(f"**{ab}** — {full}  ")
    elif kind == "ref":
        md.append(text + "\n")
    else:
        md.append(text + "\n")
with open(OUT_MD, "w", encoding="utf-8") as f:
    f.write("\n".join(md))

print("Saved", OUT_DOCX, "and", OUT_MD)
print("Body word count target ~7400; references:",
      sum(1 for k, _ in blocks if k == "ref"))
