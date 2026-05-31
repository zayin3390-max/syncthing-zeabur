#!/usr/bin/env python3
"""Reformat the verified review into the Shenzhen MSU-BIT course-assignment
template (cover page, abstract page, numbered body, References) and emit both
a .docx and a reportlab PDF preview. References are restyled to the template's
GB/T-like numeric format."""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "review_chemical_processes_climate.md"
OUT_DOCX = "review_coursework.docx"
OUT_PDF = "review_coursework_preview.pdf"

UNIVERSITY = "SHENZHEN MSU-BIT UNIVERSITY"
DEPARTMENT = "Department of Chemistry"
COURSE = "«Course: Environmental Chemistry»"
AUTHOR = "[Your Name]"
ASSIGNMENT = "Course Assignment"
INSTRUCTOR = "Dr. Shen Tianyi"
YEAR = "2025"

# Figure placeholders keyed to a unique anchor phrase in the body text. Each
# points to a real figure in a cited paper; the user downloads it and pastes it
# into the box. (No copyrighted artwork is embedded by this script.)
FIGURES = [
    dict(n=1, anchor="airborne fraction averaging 44% over 1959",
         caption="Global carbon budget: anthropogenic CO2 sources and their partitioning among atmospheric growth, the ocean sink and the land sink.",
         source="Source: Friedlingstein et al., Global Carbon Budget 2024, Earth Syst. Sci. Data, 2025, 17: 965 [2] (open access, CC BY 4.0).",
         url="https://essd.copernicus.org/articles/17/965/2025/"),
    dict(n=2, anchor="single largest source of uncertainty in total anthropogenic forcing",
         caption="Effective radiative forcing of the climate system by component over 1750-2019.",
         source="Source: IPCC AR6 WGI, Chapter 7, Fig. 7.6 (Forster et al.) [5].",
         url="https://www.ipcc.ch/report/ar6/wg1/figures/chapter-7/figure-7-6/"),
    dict(n=3, anchor="were anthropogenic",
         caption="The global methane budget: natural and anthropogenic CH4 sources and the principal sinks.",
         source="Source: Saunois et al., The Global Methane Budget 2000-2017, Earth Syst. Sci. Data, 2020, 12: 1561 [7] (open access, CC BY 4.0).",
         url="https://essd.copernicus.org/articles/12/1561/2020/"),
    dict(n=4, anchor="governed by the Revelle factor",
         caption="Change in the column inventory of anthropogenic carbon stored in the ocean.",
         source="Source: Gruber et al., Science, 2019, 363: 1193 [17] (copyright AAAS; access via publisher/institution).",
         url="https://www.science.org/doi/10.1126/science.aau5153"),
    dict(n=5, anchor="shell malformation and dissolution",
         caption="Ocean acidification: decline in surface-ocean pH and shoaling of the aragonite saturation horizon.",
         source="Source: IPCC AR6 WGI, Chapter 5, Fig. 5.21 (Canadell et al.) [6].",
         url="https://www.ipcc.ch/report/ar6/wg1/figures/chapter-5/figure-5-21/"),
    dict(n=6, anchor="nitrogen loss through denitrification",
         caption="Global distribution of declining dissolved oxygen in the open ocean and coastal waters.",
         source="Source: Breitburg et al., Science, 2018, 359: eaam7240 [23] (copyright AAAS; access via publisher/institution).",
         url="https://www.science.org/doi/10.1126/science.aam7240"),
    dict(n=7, anchor="cannot buffer anthropogenic emissions on human timescales",
         caption="The carbonate-silicate cycle and the silicate-weathering negative feedback (the long-term carbon thermostat).",
         source="Source: Penman et al., Earth-Sci. Rev., 2020, 209: 103298 [33] (copyright Elsevier; access via publisher/institution).",
         url="https://doi.org/10.1016/j.earscirev.2020.103298"),
    dict(n=8, anchor="Paris Agreement temperature targets",
         caption="Permafrost-zone carbon stocks and the permafrost carbon-climate feedback.",
         source="Source: Schuur et al., Nature, 2015, 520: 171 [46] (copyright Springer Nature; access via publisher/institution).",
         url="https://www.nature.com/articles/nature14338"),
    dict(n=9, anchor="and transient climate response of 1.8",
         caption="Individual climate feedback parameters assessed by IPCC AR6.",
         source="Source: IPCC AR6 WGI, Chapter 7, Fig. 7.10 (Forster et al.) [5].",
         url="https://www.ipcc.ch/report/ar6/wg1/figures/chapter-7/figure-7-10/"),
    dict(n=10, anchor="unresolved scientific disagreement about tipping timing",
         caption="Global map of climate tipping elements and their estimated temperature thresholds.",
         source="Source: Armstrong McKay et al., Science, 2022, 377: eabn7950 [65] (copyright AAAS; access via publisher/institution).",
         url="https://www.science.org/doi/10.1126/science.abn7950"),
]

_used_figs = set()

def fig_for(text):
    for f in FIGURES:
        if f["n"] not in _used_figs and f["anchor"] in text:
            _used_figs.add(f["n"])
            return f
    return None

def inject_callout(text, n):
    t = text.rstrip()
    if t.endswith("."):
        return t[:-1] + f" (Fig. {n})."
    return t + f" (Fig. {n})"

# ----------------------------------------------------------------- read source
with open(SRC, encoding="utf-8") as f:
    lines = f.read().split("\n")

def first_idx(pred):
    for i, l in enumerate(lines):
        if pred(l):
            return i
    return -1

title = next(l[2:].strip() for l in lines if l.startswith("# "))
i_abs = first_idx(lambda l: l.strip() == "## Abstract")
i_intro = first_idx(lambda l: l.startswith("## 1. Introduction"))
i_refs = first_idx(lambda l: l.strip() == "## References")

abstract_parts, keywords = [], ""
for l in lines[i_abs + 1:i_intro]:
    s = l.strip()
    if not s or s == "---":
        continue
    if s.startswith("**Keywords:**"):
        keywords = s.replace("**Keywords:**", "").strip()
        continue
    abstract_parts.append(s)
abstract_text = " ".join(abstract_parts)

# body blocks
body_blocks = []
for l in lines[i_intro:i_refs]:
    s = l.strip()
    if not s or s == "---":
        continue
    if s.startswith("## "):
        body_blocks.append(("h1", s[3:].strip()))
    elif s.startswith("### "):
        body_blocks.append(("h2", s[4:].strip()))
    else:
        body_blocks.append(("para", s))

# raw reference lines
raw_refs = []
for l in lines[i_refs + 1:]:
    s = l.strip()
    if s.startswith("*Prepared as"):
        break
    if re.match(r"^\[\d+\]", s):
        raw_refs.append(s)

# ----------------------------------------------------- reference restyling
INIT = re.compile(r"^(?:[A-Z]\.-?)+$")

def conv_authors(a):
    etal = "et al." in a
    a = a.replace("*et al.*", "").replace("et al.", "")
    parts = re.split(r",\s*and\s+|,\s+|\s+and\s+", a)
    out = []
    for p in parts:
        p = p.strip().strip(",")
        if not p:
            continue
        toks = p.split()
        inits = [t for t in toks if INIT.match(t)]
        sur = [t for t in toks if not INIT.match(t)]
        letters = " ".join("".join(re.findall(r"[A-Z]", t)) for t in inits)
        name = (" ".join(sur) + (" " + letters if letters else "")).strip()
        if name:
            out.append(name)
    if etal:
        out.append("et al.")
    return ", ".join(out)

def restyle_ref(s):
    m = re.match(r"^\[(\d+)\]\s+(?P<auth>.+?),\s+\"(?P<title>.+?),\"\s+(?P<tail>.+)\.$", s)
    if not m:
        # fallback: strip markdown only
        return re.sub(r"[*\"]", "", s)
    num = m.group(1)
    authors = conv_authors(m.group("auth"))
    title = m.group("title").strip()
    tail = m.group("tail").strip()

    # book chapter: "in *Book*. City: Publisher, YEAR"
    mb = re.match(r"in\s+\*(?P<book>.+?)\*\.\s*(?P<pub>.+),\s*(?P<year>\d{4})$", tail)
    if mb:
        return f"[{num}] {authors}. {title}. In: {mb.group('book')}. {mb.group('pub')}, {mb.group('year')}."

    # journal: "*Journal*, vol. X, [no. Y,] [pp. A–B | artno], YEAR"
    mj = re.match(r"\*(?P<jour>.+?)\*,\s*(?P<rest>.+)$", tail)
    if not mj:
        return re.sub(r"[*\"]", "", s)
    jour = mj.group("jour")
    rest = mj.group("rest")
    year = (re.search(r"(\d{4})\s*$", rest) or [None, ""])
    year = year.group(1) if hasattr(year, "group") else ""
    vol = re.search(r"vol\.\s*([0-9]+)", rest)
    iss = re.search(r"no\.\s*([0-9]+)", rest)
    pages = re.search(r"pp\.\s*([0-9]+(?:[–\-][0-9]+)?)", rest)
    vol = vol.group(1) if vol else ""
    iss = f"({iss.group(1)})" if iss else ""
    if pages:
        loc = pages.group(1)
    else:  # article number = token before the year that is not vol/no
        art = re.search(r",\s*([A-Za-z]?[0-9][0-9A-Za-z]*),\s*\d{4}\s*$", rest)
        loc = art.group(1) if art else ""
    locator = f"{vol}{iss}: {loc}" if loc else f"{vol}{iss}"
    return f"[{num}] {authors}. {title}. {jour}, {year}, {locator}.".replace(", , ", ", ")

refs = [restyle_ref(r) for r in raw_refs]

# ----------------------------------------------------- shared inline rendering
INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")
TOKEN = re.compile(r"\^(-?[0-9A-Za-z]+)|(CO2|CH4|N2O|SO2|NO2|CaCO3|H2O|O2|N2)")

def emit(p, text, bold=False, italic=False):
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()]); r.bold = bold; r.italic = italic
        if m.group(1) is not None:
            sup = re.match(r"(-?\d+)(.*)", m.group(1))
            if sup:
                r = p.add_run(sup.group(1)); r.font.superscript = True; r.bold = bold; r.italic = italic
                if sup.group(2):
                    r = p.add_run(sup.group(2)); r.bold = bold; r.italic = italic
            else:
                r = p.add_run(m.group(1)); r.font.superscript = True; r.bold = bold; r.italic = italic
        else:
            for ch in m.group(2):
                r = p.add_run(ch); r.bold = bold; r.italic = italic
                if ch.isdigit():
                    r.font.subscript = True
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:]); r.bold = bold; r.italic = italic

def add_runs(p, text):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            emit(p, tok[2:-2], bold=True)
        elif tok.startswith("*") and tok.endswith("*"):
            emit(p, tok[1:-1], italic=True)
        else:
            emit(p, tok)

# =====================================================================  DOCX
doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

def para(align=None, before=0, after=0, spacing=WD_LINE_SPACING.ONE_POINT_FIVE):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = spacing
    return p

def run(p, text, size=12, bold=False, italic=False):
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    return r

def blank(n=1):
    for _ in range(n):
        para(spacing=WD_LINE_SPACING.SINGLE)

def hrule(p):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "10"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(k), v)
    pbdr.append(bottom); pPr.append(pbdr)

C = WD_ALIGN_PARAGRAPH.CENTER
R = WD_ALIGN_PARAGRAPH.RIGHT
J = WD_ALIGN_PARAGRAPH.JUSTIFY

# --- cover page
blank(1)
run(para(C, after=4, spacing=WD_LINE_SPACING.SINGLE), UNIVERSITY, 14, bold=True)
hrule(para(spacing=WD_LINE_SPACING.SINGLE))
blank(1)
run(para(C, after=6, spacing=WD_LINE_SPACING.SINGLE), DEPARTMENT, 16, bold=True)
run(para(C, spacing=WD_LINE_SPACING.SINGLE), COURSE, 13)
blank(3)
run(para(C, after=4, spacing=WD_LINE_SPACING.SINGLE), AUTHOR, 13)
pt = para(C, after=4, spacing=WD_LINE_SPACING.SINGLE)
run(pt, title, 18, bold=True)
run(para(C, spacing=WD_LINE_SPACING.SINGLE), ASSIGNMENT, 12)
blank(2)
run(para(R, after=2, spacing=WD_LINE_SPACING.SINGLE), "Instructor:", 12)
run(para(R, spacing=WD_LINE_SPACING.SINGLE), INSTRUCTOR, 12)
blank(8)
run(para(C, spacing=WD_LINE_SPACING.SINGLE), YEAR, 12)
doc.add_page_break()

# --- abstract page
run(para(C, after=10, spacing=WD_LINE_SPACING.SINGLE), "Abstract", 13, bold=True)
ap = para(J); ap.paragraph_format.first_line_indent = Pt(18)
emit(ap, abstract_text)
blank(1)
kp = para(J)
run(kp, "Keywords: ", 12, bold=True)
emit(kp, keywords)
doc.add_page_break()

def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "0"), ("w:color", "999999")):
            e.set(qn(k), v)
        borders.append(e)
    tcPr.append(borders)

def add_figure_docx(f):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0); cell.width = Inches(6.0)
    set_cell_border(cell)
    cp = cell.paragraphs[0]; cp.alignment = C
    cp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    cp.paragraph_format.space_before = Pt(46); cp.paragraph_format.space_after = Pt(46)
    r = cp.add_run(f"[ Fig. {f['n']} — paste figure here ]")
    r.italic = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    cap = para(J, before=2, after=10, spacing=WD_LINE_SPACING.SINGLE)
    run(cap, f"Fig. {f['n']}. ", 9.5, bold=True)
    run(cap, f["caption"] + " ", 9.5)
    run(cap, f["source"], 9.5, italic=True)
    run(cap, "  Download: " + f["url"], 9.5)

# --- body
for kind, text in body_blocks:
    if kind == "h1":
        p = para(before=12, after=4, spacing=WD_LINE_SPACING.SINGLE)
        run(p, "", 13); add_runs(p, text)
        for r in p.runs:
            r.bold = True; r.font.size = Pt(13)
    elif kind == "h2":
        p = para(before=8, after=2, spacing=WD_LINE_SPACING.SINGLE)
        add_runs(p, text)
        for r in p.runs:
            r.bold = True; r.font.size = Pt(12)
    else:
        fig = fig_for(text)
        p = para(J)
        p.paragraph_format.first_line_indent = Pt(18)
        add_runs(p, inject_callout(text, fig["n"]) if fig else text)
        if fig:
            add_figure_docx(fig)

# --- references
doc.add_page_break()
run(para(C, after=10, spacing=WD_LINE_SPACING.SINGLE), "References", 13, bold=True)
for r in refs:
    p = para(J, after=4, spacing=WD_LINE_SPACING.SINGLE)
    p.paragraph_format.left_indent = Pt(22)
    p.paragraph_format.first_line_indent = Pt(-22)
    add_runs(p, r)

doc.save(OUT_DOCX)

# =====================================================================  PDF
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame,
                                Paragraph, Spacer, PageBreak,
                                HRFlowable, Table, TableStyle)
from reportlab.lib import colors

def pesc(t):
    t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"<i>\1</i>", t)
    t = re.sub(r"\^(-?\d+)", r"<super>\1</super>", t)
    for f in ("CO2", "CH4", "N2O", "SO2", "NO2", "CaCO3", "H2O", "O2", "N2"):
        t = t.replace(f, re.sub(r"(\d+)", r"<sub>\1</sub>", f))
    return t

S = {
    "uni": ParagraphStyle("uni", fontName="Times-Bold", fontSize=14, leading=18, alignment=TA_CENTER),
    "dept": ParagraphStyle("dept", fontName="Times-Bold", fontSize=16, leading=20, alignment=TA_CENTER, spaceAfter=6),
    "course": ParagraphStyle("course", fontName="Times-Roman", fontSize=13, leading=17, alignment=TA_CENTER),
    "author": ParagraphStyle("author", fontName="Times-Roman", fontSize=13, leading=17, alignment=TA_CENTER, spaceAfter=4),
    "ptitle": ParagraphStyle("ptitle", fontName="Times-Bold", fontSize=18, leading=23, alignment=TA_CENTER, spaceAfter=4),
    "assign": ParagraphStyle("assign", fontName="Times-Roman", fontSize=12, leading=16, alignment=TA_CENTER),
    "right": ParagraphStyle("right", fontName="Times-Roman", fontSize=12, leading=16, alignment=TA_RIGHT),
    "yc": ParagraphStyle("yc", fontName="Times-Roman", fontSize=12, leading=16, alignment=TA_CENTER),
    "head": ParagraphStyle("head", fontName="Times-Bold", fontSize=13, leading=17, alignment=TA_CENTER, spaceAfter=10),
    "h1": ParagraphStyle("h1", fontName="Times-Bold", fontSize=13, leading=17, spaceBefore=12, spaceAfter=4),
    "h2": ParagraphStyle("h2", fontName="Times-Bold", fontSize=12, leading=15, spaceBefore=8, spaceAfter=2),
    "body": ParagraphStyle("body", fontName="Times-Roman", fontSize=12, leading=18, alignment=TA_JUSTIFY, firstLineIndent=18),
    "kw": ParagraphStyle("kw", fontName="Times-Roman", fontSize=12, leading=18, alignment=TA_JUSTIFY),
    "ref": ParagraphStyle("ref", fontName="Times-Roman", fontSize=11, leading=15, leftIndent=22, firstLineIndent=-22, spaceAfter=4, alignment=TA_JUSTIFY),
    "phbox": ParagraphStyle("phbox", fontName="Times-Italic", fontSize=10, leading=14, alignment=TA_CENTER, textColor=colors.grey),
    "cap": ParagraphStyle("cap", fontName="Times-Roman", fontSize=9, leading=12, alignment=TA_JUSTIFY, spaceBefore=3, spaceAfter=8),
}

def fig_flowables(f, width):
    box = Table([[Paragraph(f"[ Fig. {f['n']} — paste figure here ]", S["phbox"])]],
                colWidths=[width], rowHeights=[150])
    box.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, colors.grey),
                             ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    cap = Paragraph(f"<b>Fig. {f['n']}.</b> {pesc(f['caption'])} <i>{pesc(f['source'])}</i> "
                    f"Download: {f['url']}", S["cap"])
    return [Spacer(1, 6), box, cap]

story = [Spacer(1, 20)]
story.append(Paragraph(pesc(UNIVERSITY), S["uni"]))
story.append(Spacer(1, 4))
story.append(HRFlowable(width="100%", thickness=1, color="black"))
story.append(Spacer(1, 18))
story.append(Paragraph(pesc(DEPARTMENT), S["dept"]))
story.append(Paragraph(pesc(COURSE), S["course"]))
story.append(Spacer(1, 60))
story.append(Paragraph(pesc(AUTHOR), S["author"]))
story.append(Paragraph(pesc(title), S["ptitle"]))
story.append(Paragraph(pesc(ASSIGNMENT), S["assign"]))
story.append(Spacer(1, 40))
story.append(Paragraph("Instructor:", S["right"]))
story.append(Paragraph(pesc(INSTRUCTOR), S["right"]))
story.append(Spacer(1, 150))
story.append(Paragraph(YEAR, S["yc"]))
story.append(PageBreak())

story.append(Paragraph("Abstract", S["head"]))
story.append(Paragraph(pesc(abstract_text), S["body"]))
story.append(Spacer(1, 10))
story.append(Paragraph("<b>Keywords:</b> " + pesc(keywords), S["kw"]))
story.append(PageBreak())

PDF_W = A4[0] - 50 * mm
_used_figs.clear()
for kind, text in body_blocks:
    if kind in ("h1", "h2"):
        story.append(Paragraph(pesc(text), S[kind]))
    else:
        fig = fig_for(text)
        story.append(Paragraph(pesc(inject_callout(text, fig["n"]) if fig else text), S["body"]))
        if fig:
            for fl in fig_flowables(fig, PDF_W):
                story.append(fl)

story.append(PageBreak())
story.append(Paragraph("References", S["head"]))
for r in refs:
    story.append(Paragraph(pesc(r), S["ref"]))

frame = Frame(25 * mm, 20 * mm, A4[0] - 50 * mm, A4[1] - 40 * mm, id="f")
BaseDocTemplate(OUT_PDF, pagesize=A4,
                pageTemplates=[PageTemplate(id="all", frames=[frame])]).build(story)

print("Saved", OUT_DOCX, "and", OUT_PDF)
print("references restyled:", len(refs))
