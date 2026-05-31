#!/usr/bin/env python3
"""Convert the review Markdown into a formatted .docx (no external services)."""
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = "review_chemical_processes_climate.md"
OUT = "review_chemical_processes_climate.docx"

doc = Document()

# Base style
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)

INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")

def add_runs(paragraph, text):
    """Add text to paragraph, honoring **bold** and *italic* markers."""
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = paragraph.add_run(tok[2:-2]); run.bold = True
        elif tok.startswith("*") and tok.endswith("*"):
            run = paragraph.add_run(tok[1:-1]); run.italic = True
        else:
            paragraph.add_run(tok)

with open(SRC, encoding="utf-8") as f:
    lines = f.readlines()

in_refs = False
for raw in lines:
    line = raw.rstrip("\n")
    stripped = line.strip()

    if stripped == "" :
        continue
    if stripped == "---":
        continue

    if line.startswith("# "):
        p = doc.add_heading(line[2:].strip(), level=0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        continue
    if line.startswith("## "):
        heading = line[3:].strip()
        doc.add_heading(heading, level=1)
        in_refs = heading.lower().startswith("references")
        continue
    if line.startswith("### "):
        doc.add_heading(line[4:].strip(), level=2)
        continue

    # Reference entries: keep as hanging-ish single-spaced list
    if in_refs and re.match(r"^\[\d+\]", stripped):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        add_runs(p, stripped)
        continue

    # Closing italic note line (whole-line italics)
    if stripped.startswith("*") and stripped.endswith("*") and stripped.count("*") == 2:
        p = doc.add_paragraph()
        run = p.add_run(stripped[1:-1]); run.italic = True; run.font.size = Pt(9)
        continue

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    add_runs(p, stripped)

doc.save(OUT)
print("Saved", OUT)
